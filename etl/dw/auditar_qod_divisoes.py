# requirements:
# openpyxl, psycopg2-binary
"""Acrescenta ao anexo da auditoria do QOD (auditar_qod.py) o que depende de casar o
QOD com as divisoes do organograma -- as mesmas regras da carga (carga_core.py,
carregar_divisoes_qod). NAO corrige nada.

Abas acrescentadas:
  Divisoes acima do previsto   lotacao (CPPA) maior que o previsto do quadro (DPlaG)
  Secoes sem no no quadro      secao do catalogo de funcoes que nao existe na arvore do
                               quadro da unidade (DPlaG: catalogo x quadro)
  Funcao em outra unidade      ocupante lotado, no Peculio, em unidade diferente da OME da
                               funcao (CPPA/DPlaG)
  Codigos sem funcao           codigo da QOD_CONSOLIDADA que nao esta no catalogo (CPPA)

Uso:
    DATABASE_URL_DW_ETL=... python auditar_qod_divisoes.py anexo.xlsx
"""
import os
import sys
from collections import Counter, defaultdict
from pathlib import Path

import openpyxl
import psycopg2

sys.path.insert(0, str(Path(__file__).resolve().parent))
import carga_core as C  # noqa: E402


def main(caminho):
    conn = psycopg2.connect(os.environ["DATABASE_URL_DW_ETL"])
    conn.set_session(readonly=True)
    cur = conn.cursor()

    previsto, avisos, herancas = C.previsto_das_divisoes(cur)
    previsto_por = {(s, n): p for s, itens in previsto.items() for n, p, _ in itens}
    divisao_da_funcao, lotacao, _ = C.lotacao_das_divisoes(cur)

    wb = openpyxl.load_workbook(caminho)
    for nome in ("Divisoes acima do previsto", "Secoes sem no no quadro", "Funcao em outra unidade", "Codigos sem funcao"):
        if nome in wb.sheetnames:
            del wb[nome]

    # 1. lotacao acima do previsto
    linhas = []
    for (sigla, divisao), mats in lotacao.items():
        prev = previsto_por.get((sigla, divisao), 0)
        if len(mats) > prev:
            linhas.append((sigla, divisao, prev, len(mats), len(mats) - prev))
    linhas.sort(key=lambda l: -l[4])
    ws = wb.create_sheet("Divisoes acima do previsto")
    ws.append(["unidade", "divisao", "previsto (quadro)", "lotados (QOD)", "diferenca"])
    for l in linhas:
        ws.append(list(l))

    # 2. secoes do catalogo que nao existem na arvore do quadro da unidade
    cur.execute("SELECT valor_origem FROM stg.de_para_ome WHERE sistema = 'organograma'")
    siglas = {s for (s,) in cur.fetchall()}
    cur.execute("SELECT sigla, nome FROM stg.divisao_organograma")
    for s, _ in cur.fetchall():
        siglas.add(s)
    cur.execute("""
        SELECT f.id_funcao, f.ome_qod, coalesce(f.secao, ''), coalesce(f.funcao, ''),
               (SELECT count(*) FROM core.qod_ocupacao o WHERE o.id_funcao = f.id_funcao)
        FROM core.qod_funcao f WHERE f.ativo
    """)
    sem_no = Counter()
    secao_ids = defaultdict(list)
    for id_funcao, ome, secao, funcao, n_ocup in cur.fetchall():
        sigla = C.unidade_pai_do_qod((ome or "").strip(), siglas)
        if (ome or "").strip() != sigla:  # secao operacional (SB): casa pela OME, nao pela arvore
            continue
        arvore = herancas.get(sigla)
        if arvore is None or not secao.strip():
            continue
        if C._norm_quadro(secao) not in arvore:
            sem_no[(sigla, secao.strip())] += 1
            secao_ids[(sigla, secao.strip())].append(id_funcao)
    ws = wb.create_sheet("Secoes sem no no quadro")
    ws.append(["unidade", "secao (catalogo de funcoes)", "funcoes", "exemplo de ID_FUNCAO"])
    for (sigla, secao), n in sorted(sem_no.items(), key=lambda kv: (-kv[1], kv[0])):
        ws.append([sigla, secao, n, secao_ids[(sigla, secao)][0]])

    # 3. ocupante lotado em unidade diferente da OME da funcao
    cur.execute("SELECT matricula_normalizada, sigla FROM marts.efetivo_nominal")
    unidade_do_militar = dict(cur.fetchall())
    cur.execute("""
        SELECT o.linha, o.matricula, o.id_funcao, f.ome_qod, f.secao, f.funcao
        FROM core.qod_ocupacao o JOIN core.qod_funcao f ON f.id_funcao = o.id_funcao
    """)
    ws = wb.create_sheet("Funcao em outra unidade")
    ws.append(["linha (QOD_CONSOLIDADA)", "matricula", "ID_FUNCAO", "OME da funcao", "unidade no Peculio", "secao", "funcao"])
    contagem = Counter()
    for linha, matricula, id_funcao, ome, secao, funcao in cur.fetchall():
        if matricula not in unidade_do_militar:
            continue
        sigla_funcao = C.unidade_pai_do_qod((ome or "").strip(), siglas)
        if sigla_funcao != unidade_do_militar[matricula]:
            contagem[(sigla_funcao, unidade_do_militar[matricula])] += 1
            ws.append([linha, matricula, id_funcao, ome, unidade_do_militar[matricula], secao, funcao])
    ws_res = wb["Resumo"] if "Resumo" in wb.sheetnames else None

    # 4. codigos da QOD_CONSOLIDADA sem funcao no catalogo
    cur.execute("""
        SELECT o.linha, o.matricula, o.unidade_qod, o.id_funcao
        FROM core.qod_ocupacao o LEFT JOIN core.qod_funcao f ON f.id_funcao = o.id_funcao
        WHERE o.id_funcao IS NOT NULL AND f.id_funcao IS NULL ORDER BY o.linha
    """)
    ws = wb.create_sheet("Codigos sem funcao")
    ws.append(["linha (QOD_CONSOLIDADA)", "matricula", "unidade (QOD)", "codigo"])
    n_sem = 0
    for l in cur.fetchall():
        ws.append(list(l))
        n_sem += 1

    wb.save(caminho)
    print(f"divisoes acima do previsto: {len(linhas)} ({sum(l[4] for l in linhas)} pessoas)")
    print(f"secoes do catalogo sem no no quadro: {len(sem_no)} secoes, {sum(sem_no.values())} funcoes")
    print(f"ocupantes em unidade diferente da OME da funcao: {sum(contagem.values())}")
    print(f"codigos sem funcao no catalogo: {n_sem}")


if __name__ == "__main__":
    main(sys.argv[1])
