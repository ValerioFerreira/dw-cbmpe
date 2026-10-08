"""Relatorio (xlsx) das ocupacoes da QOD_CONSOLIDADA que ficaram sem funcao depois da traducao
dos codigos antigos para o catalogo ID_ATUALIZADO (carga_core.traduzir_codigos_qod).

Para a DPlaG (dona do catalogo de funcoes) e a CPPA (dona de quem ocupa cada uma):
  - Ocupacoes sem funcao: o CODIGO escrito na QOD_CONSOLIDADA aponta para uma funcao do
    catalogo antigo que nao tem par no novo. Mostra a funcao antiga e, quando ha, candidatas
    do catalogo novo (mesma funcao, mesma unidade-mae) que ficaram sem ocupante traduzido.
  - Funcoes antigas sem par: o que a DPlaG precisa dizer se deixou de existir ou mudou de nome.
  - Funcoes novas sem par: criadas no ID_ATUALIZADO sem correspondente no catalogo antigo.

Somente leitura no DW. Uso (raiz do repositorio, com .env.local):
    python etl/dw/relatorio_qod_sem_funcao.py            # grava docs/relatorios/QOD_ocupacoes_sem_funcao.xlsx
"""
import re
import sys
from collections import defaultdict
from datetime import date
from pathlib import Path

import psycopg2
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

RAIZ = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(RAIZ / "etl" / "dw"))
import carga_core as cc  # noqa: E402

SAIDA = RAIZ / "docs" / "relatorios" / "QOD_ocupacoes_sem_funcao.xlsx"


def conectar():
    env = {}
    for linha in (RAIZ / ".env.local").read_text(encoding="utf-8").splitlines():
        m = re.match(r"^([A-Z_0-9]+)=(.*)$", linha.strip())
        if m:
            env[m.group(1)] = m.group(2).strip().strip('"')
    return psycopg2.connect(env["DATABASE_URL_DW"].rpartition("/")[0] + "/dw")


def folha(wb, titulo, cabecalho, linhas, larguras):
    ws = wb.create_sheet(titulo)
    ws.append(cabecalho)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="8B0000")
        c.alignment = Alignment(wrap_text=True, vertical="center")
    for linha in linhas:
        ws.append(list(linha))
    for i, w in enumerate(larguras, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    return ws


def main():
    c = conectar()
    cur = c.cursor()
    cur.execute("""SELECT trim(col_01_id_funcao), col_05_ome, col_07_secao, col_08_funcao
                   FROM raw_qod.qod_funcoes WHERE nullif(trim(col_01_id_funcao), '') IS NOT NULL""")
    antigo = cur.fetchall()
    cur.execute("SELECT id_funcao, ome_qod, secao, funcao FROM core.qod_funcao")
    novo = cur.fetchall()
    mapa = cc.alinhar_catalogos(antigo, novo)
    info_antigo = {i: (o, s, f) for i, o, s, f in antigo}
    info_novo = {i: (o, s, f) for i, o, s, f in novo}
    novos_com_par = set(mapa.values())

    # funcoes novas sem par, indexadas por (unidade-mae, funcao) para sugerir candidatas
    candidatas = defaultdict(list)
    for i, o, s, f in novo:
        if i not in novos_com_par:
            candidatas[(cc._ome_pai_norm(o), cc._norm_lot(f))].append((i, o, s, f))

    cur.execute("""
        SELECT q.linha, q.matricula, q.posto_grad, q.nome_qod, q.unidade_qod, trim(r.col_13_codigo), q.tipo, q.funcao_de
        FROM core.qod_ocupacao q JOIN raw_qod.qod_consolidada r ON r.linha = q.linha
        ORDER BY q.unidade_qod, q.posto_grad, q.nome_qod""")
    ocupacoes = cur.fetchall()

    sem_funcao, codigo_inexistente = [], []
    for linha, mat, posto, nome, unidade, codigo, tipo, funcao_de in ocupacoes:
        if not codigo or codigo in mapa:
            continue
        ant = info_antigo.get(codigo)
        if not ant:
            codigo_inexistente.append((linha, mat, posto, nome, unidade, codigo, tipo, funcao_de))
            continue
        cand = candidatas.get((cc._ome_pai_norm(ant[0]), cc._norm_lot(ant[2])), [])
        sugestao = "; ".join(f"{i} ({o} / {s})" for i, o, s, _ in cand[:3])
        sem_funcao.append((linha, mat, posto, nome, unidade, codigo, ant[0], ant[1], ant[2], tipo, funcao_de,
                           sugestao or "nenhuma funcao nova com o mesmo nome na mesma unidade"))

    ocupados_antigos = {o[5] for o in ocupacoes if o[5]}
    antigas_sem_par = [(i, *info_antigo[i], "sim" if i in ocupados_antigos else "nao")
                       for i in sorted(info_antigo, key=lambda x: cc._numero_id(x)) if i not in mapa]
    novas_sem_par = [(i, o, s, f) for i, o, s, f in sorted(novo, key=lambda x: cc._numero_id(x[0]))
                     if i not in novos_com_par]

    wb = Workbook()
    ws = wb.active
    ws.title = "Resumo"
    resumo = [
        ("Relatorio QOD: ocupacoes sem funcao", ""),
        ("Gerado em", date.today().strftime("%d/%m/%Y")),
        ("", ""),
        ("Contexto", "A coluna CODIGO da QOD_CONSOLIDADA segue a numeracao antiga das funcoes (QOD_FUNCOES, agosto); "
                     "o catalogo oficial passou a ser o ID_ATUALIZADO. O DW traduz um para o outro alinhando por "
                     "(OME, secao, funcao) e, depois, por (OME, funcao)."),
        ("Ocupacoes na QOD_CONSOLIDADA", len(ocupacoes)),
        ("Funcoes no catalogo antigo / novo", f"{len(antigo)} / {len(novo)}"),
        ("Codigos antigos traduzidos", len(mapa)),
        ("Ocupacoes sem funcao (codigo antigo sem par)", len(sem_funcao)),
        ("Ocupacoes com codigo que nem existe no catalogo antigo", len(codigo_inexistente)),
        ("", ""),
        ("DPlaG (funcoes)", "Abas 'Funcoes antigas sem par' e 'Funcoes novas sem par': a funcao deixou de existir, "
                            "mudou de nome/secao ou foi duplicada? Se mudou, informar o ID novo correspondente."),
        ("CPPA (ocupantes)", "Aba 'Ocupacoes sem funcao': conferir o codigo escrito na linha; se a funcao existe com outro "
                             "ID, informar. Aba 'Codigo inexistente': codigos que nao constam em nenhum catalogo."),
        ("Enquanto isso", "Esses militares aparecem lotados na unidade escrita na linha, sem funcao/secao/divisao."),
    ]
    for linha in resumo:
        ws.append(list(linha))
    ws["A1"].font = Font(bold=True, size=14)
    ws.column_dimensions["A"].width = 52
    ws.column_dimensions["B"].width = 110
    for row in ws.iter_rows():
        for cel in row:
            cel.alignment = Alignment(wrap_text=True, vertical="top")

    folha(wb, "Ocupacoes sem funcao",
          ["Linha QOD_CONSOLIDADA", "Matricula", "Posto/Grad", "Nome", "Unidade escrita na linha", "Codigo escrito",
           "OME da funcao antiga", "Secao antiga", "Funcao antiga", "Tipo", "Funcao de", "Candidatas no catalogo novo"],
          sem_funcao, [10, 12, 10, 34, 24, 10, 20, 38, 34, 10, 22, 60])
    folha(wb, "Codigo inexistente",
          ["Linha QOD_CONSOLIDADA", "Matricula", "Posto/Grad", "Nome", "Unidade escrita na linha", "Codigo escrito",
           "Tipo", "Funcao de"], codigo_inexistente, [10, 12, 10, 34, 24, 12, 10, 22])
    folha(wb, "Funcoes antigas sem par",
          ["ID antigo", "OME", "Secao", "Funcao", "Tem ocupante na QOD_CONSOLIDADA?"],
          antigas_sem_par, [10, 20, 44, 40, 18])
    folha(wb, "Funcoes novas sem par", ["ID novo (ID_ATUALIZADO)", "OME", "Secao", "Funcao"],
          novas_sem_par, [16, 20, 44, 40])
    SAIDA.parent.mkdir(parents=True, exist_ok=True)
    wb.save(SAIDA)
    print(f"{SAIDA}\n  ocupacoes sem funcao: {len(sem_funcao)}; codigo inexistente: {len(codigo_inexistente)}; "
          f"funcoes antigas sem par: {len(antigas_sem_par)} (com ocupante: {sum(1 for a in antigas_sem_par if a[4] == 'sim')}); "
          f"novas sem par: {len(novas_sem_par)}")
    c.rollback()


if __name__ == "__main__":
    main()
