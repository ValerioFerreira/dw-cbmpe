# requirements:
# openpyxl
"""Extração pontual do Efetivo Previsto (QOD) a partir da planilha oficial
compartilhada pelo usuário, para popular src/data/real/efetivo_previsto.json.

Uso:
    python extract_previsto.py "C:\\caminho\\para\\08 - QOD - AGOSTO - 2026.xlsx"

Fonte: aba "GERAL 001" (QUADRO DE ORGANIZAÇÃO E DISTRIBUIÇÃO / DISTRIBUIÇÃO
GERAL DE EFETIVO) — coluna DISCRIMINAÇÃO (unidade) x coluna TOTAL (previsto
próprio daquela unidade, sem agregar subordinadas — a agregação acontece em
organograma.js, igual efetivo/viaturas/financas reais).

Decisões confirmadas com o usuário (ver DW_PLANEJAMENTO.md, 2026-08-11):
- Usar mesmo com o QOD pendente de homologação em BGE.
- Ignorar as abas "mergulho"/"sede"/"destacada" (estudo de reorganização à
  parte, fora do escopo do previsto vigente).
- CPPA soma dentro de DGP (mesma regra do Efetivo atual/Viaturas/Finanças).
- DGA não aparece no QOD (unidade planejada) — fica com previsto=0.

Nenhum dado individual sai daqui — só o total previsto por unidade
(agregado), mesma minimização das outras extrações.
"""
import json
import sys
from datetime import datetime, timezone
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[2]
MAPEAMENTO_PATH = Path(__file__).resolve().parent / "mapeamento_previsto.json"
OUTPUT_PATH = ROOT / "src" / "data" / "real" / "efetivo_previsto.json"

COL_DISCRIMINACAO = 1
COL_TOTAL = 21
TOTAL_GERAL_ESPERADO = 5249


def normalizar(s):
    return " ".join(str(s).split())


def main():
    if len(sys.argv) < 2:
        raise SystemExit("Uso: python extract_previsto.py <caminho do xlsx do QOD>")
    xlsx_path = Path(sys.argv[1])

    mapeamento_bruto = json.loads(MAPEAMENTO_PATH.read_text(encoding="utf-8"))
    mapeamento = {
        normalizar(k): v for k, v in mapeamento_bruto.items() if not k.startswith("_")
    }

    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    ws = wb["GERAL 001"]

    por_unidade = {}
    nao_mapeados = []
    total_lido = 0

    for row in ws.iter_rows(min_row=6, max_row=ws.max_row):
        discriminacao = row[COL_DISCRIMINACAO - 1].value
        total = row[COL_TOTAL - 1].value
        if discriminacao is None:
            continue
        chave = normalizar(discriminacao)
        if chave == "TOTAL GERAL":
            break
        if total is None:
            continue
        total = int(float(total))

        if chave not in mapeamento:
            nao_mapeados.append(chave)
            continue

        total_lido += total
        sigla = mapeamento[chave]
        if sigla is None:
            # CPPA (e futuras exceções do tipo) — soma na unidade-mãe.
            sigla = "DGP"
        por_unidade[sigla] = por_unidade.get(sigla, 0) + total

    if nao_mapeados:
        raise SystemExit(
            "Linhas da aba GERAL 001 sem mapeamento em mapeamento_previsto.json: "
            f"{nao_mapeados}. Atualize o mapeamento antes de gerar o JSON."
        )

    if total_lido != TOTAL_GERAL_ESPERADO:
        raise SystemExit(
            f"Soma lida ({total_lido}) não bate com o TOTAL GERAL esperado "
            f"({TOTAL_GERAL_ESPERADO}). Não gravando o JSON — confira a aba "
            "GERAL 001 e o mapeamento antes de tentar de novo."
        )

    saida = {
        "fonte": "QOD (08 - QOD - AGOSTO - 2026.xlsx, aba GERAL 001) — pendente de homologação/publicação em BGE",
        "gerado_em": datetime.now(timezone.utc).isoformat(),
        "total_previsto": total_lido,
        "por_unidade": por_unidade,
    }

    OUTPUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    OUTPUT_PATH.write_text(
        json.dumps(saida, ensure_ascii=False, indent=2, sort_keys=True),
        encoding="utf-8",
    )
    print(f"Efetivo previsto gravado em {OUTPUT_PATH} — total: {total_lido}")


if __name__ == "__main__":
    main()
